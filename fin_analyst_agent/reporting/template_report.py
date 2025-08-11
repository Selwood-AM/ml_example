from __future__ import annotations

from typing import Any, List, Optional, Tuple
import os
import shutil
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple


def _normalize_name(name: str) -> str:
    return name.replace(" ", "_")


def _resolve_defined_name_cell(wb, name: str) -> Optional[Tuple[str, str]]:
    # Returns (sheet_name, cell_ref) if found
    dn = None
    for defined in wb.defined_names.definedName:
        if defined.name == name:
            dn = defined
            break
    if dn is None:
        return None
    # DefinedName.destinations yields tuples of (sheet_name, cell_range)
    dests = list(dn.destinations)
    if not dests:
        return None
    sheet_name, ref = dests[0]
    # We only handle single-cell references here
    if ":" in ref:
        # range provided; use top-left
        ref = ref.split(":", 1)[0]
    return sheet_name, ref


class TemplateExcelReportBuilder:
    def __init__(self, template_path: str, output_path: str) -> None:
        self.template_path = template_path
        self.output_path = output_path

    def _write_named_values(self, wb, values: dict[str, Any]) -> None:
        for key, value in values.items():
            candidates = [key, _normalize_name(key)]
            for cand in candidates:
                loc = _resolve_defined_name_cell(wb, cand)
                if loc is None:
                    continue
                sheet_name, cell = loc
                ws = wb[sheet_name]
                ws[cell] = value
                break
            # silently skip if not found

    def _write_table(self, wb, sheet_name: str, start_cell: str, headers: List[str], rows: List[List[Any]]) -> None:
        ws = wb[sheet_name]
        start_row, start_col = coordinate_to_tuple(start_cell)
        # headers
        for idx, h in enumerate(headers):
            ws.cell(row=start_row, column=start_col + idx, value=h)
        # rows
        for r_idx, row in enumerate(rows):
            for c_idx, val in enumerate(row):
                ws.cell(row=start_row + 1 + r_idx, column=start_col + c_idx, value=val)

    def build(
        self,
        assumptions: dict,
        fcfs: List[float],
        pvs: List[float],
        tv: float,
        tv_pv: float,
        ev: float,
        per_share: float | None,
        dcf_anchor_name: str = "DCF_Start",
        dcf_sheet_fallback: Optional[str] = None,
        dcf_cell_fallback: Optional[str] = None,
    ) -> str:
        os.makedirs(os.path.dirname(self.output_path) or ".", exist_ok=True)
        shutil.copyfile(self.template_path, self.output_path)
        wb = load_workbook(self.output_path)

        # Scalars: assumptions and outputs
        scalar_values = {
            **assumptions,
            "Terminal Value": tv,
            "Terminal Value PV": tv_pv,
            "Enterprise Value": ev,
            "Intrinsic Value/Share": per_share,
        }
        self._write_named_values(wb, scalar_values)

        # DCF table placement via named anchor (preferred) or fallback coordinates
        anchor = _resolve_defined_name_cell(wb, dcf_anchor_name)
        if anchor is None and dcf_sheet_fallback and dcf_cell_fallback:
            sheet_name, cell = dcf_sheet_fallback, dcf_cell_fallback
        elif anchor is not None:
            sheet_name, cell = anchor
        else:
            # default to sheet named "DCF" cell A1 if nothing given
            sheet_name, cell = ("DCF", "A1")
            if sheet_name not in wb.sheetnames:
                wb.create_sheet(sheet_name)
        headers = ["Year", "FCF", "PV of FCF"]
        rows = [[i + 1, fcfs[i], pvs[i]] for i in range(len(fcfs))]
        self._write_table(wb, sheet_name, cell, headers, rows)

        wb.save(self.output_path)
        return self.output_path